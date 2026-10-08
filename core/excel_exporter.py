# -*- coding: utf-8 -*-
import os
import shutil
from typing import List
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.worksheet.pagebreak import Break
from openpyxl.utils import get_column_letter

from core.models import PictureItem
from core.image_processor import process_picture

class ExcelExporter:
    """기성사진대지26.08.xlsx 양식과 100% 호환되는 사진대지 엑셀 생성기"""
    def __init__(self, default_project_name: str = "공사명 미지정"):
        self.default_project_name = default_project_name

    def export(self, items: List[PictureItem], output_path: str, project_name: str = None, table_config: dict = None) -> bool:
        if not items:
            return False

        proj_name = project_name or self.default_project_name
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "공사사진첩"

        # 1. 인쇄 및 페이지 설정 (A4 세로, 여백, 반복행, 96% 맞춤 배율)
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
        ws.page_setup.scale = 96
        ws.print_title_rows = "$1:$2"
        ws.page_margins.top = 0.748
        ws.page_margins.bottom = 0.748
        ws.page_margins.left = 0.63
        ws.page_margins.right = 0.55

        # 2. 열 너비 설정 (A ~ I)
        col_widths = {
            "A": 5.63, "B": 8.38, "C": 10.63, "D": 12.38, 
            "E": 10.63, "F": 5.88, "G": 10.13, "H": 9.88, "I": 8.25
        }
        for col_letter, width in col_widths.items():
            ws.column_dimensions[col_letter].width = width

        # 스타일 정의
        font_title = Font(name="굴림체", size=24, bold=True)
        font_proj = Font(name="굴림체", size=11, bold=False)
        font_label = Font(name="굴림체", size=11, bold=True)
        font_data = Font(name="굴림체", size=11, bold=False)
        
        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")
        
        thin_border_side = Side(style="thin", color="000000")
        medium_border_side = Side(style="medium", color="000000")
        
        box_border = Border(
            left=thin_border_side, right=thin_border_side,
            top=thin_border_side, bottom=thin_border_side
        )

        # 3. 1행(제목) & 2행(공사명) 작성
        ws.row_dimensions[1].height = 33.75
        ws.merge_cells("A1:I1")
        ws["A1"] = "사  진  대  지"
        ws["A1"].font = font_title
        ws["A1"].alignment = align_center

        ws.row_dimensions[2].height = 22.5
        ws["A2"] = f"공사명 : {proj_name}"
        ws["A2"].font = font_proj
        ws["A2"].alignment = align_left

        # 임시 이미지 폴더
        temp_dir = os.path.join(os.path.dirname(output_path), "_temp_excel_imgs")
        os.makedirs(temp_dir, exist_ok=True)

        try:
            # 4. 사진 및 표 슬롯 생성 (A4 1페이지 2장 규격 완성을 위해 짝수 슬롯으로 맞춤)
            total_slots = len(items)
            if total_slots % 2 != 0:
                total_slots += 1  # 홀수 장일 경우 마지막 페이지 하단 빈 슬롯 자동 완성

            # table_config 분석 (1~2개 선택 시 1줄, 3~4개 선택 시 2줄)
            default_fields = [
                {"key": "work_type", "label": "공    종", "enabled": True},
                {"key": "content", "label": "내       용", "enabled": True},
                {"key": "location", "label": "위    치", "enabled": True},
                {"key": "date_str", "label": "촬 영 일 자", "enabled": True},
            ]
            cfg_fields = (table_config.get("fields") if table_config else None) or default_fields
            active_fields = [f for f in cfg_fields if f.get("enabled", True)]
            if not active_fields:
                active_fields = default_fields
                
            field_count = len(active_fields)
            row_count = 1 if field_count <= 2 else 2

            def _get_val(it, key: str) -> str:
                if not it: return ""
                if key == "work_type": return it.work_type or ""
                if key == "content": return it.content or ""
                if key == "location": return it.location or ""
                if key == "date_str": return it.date_str or ""
                if key == "custom1": return getattr(it, "custom1", "") or ""
                if key == "custom2": return getattr(it, "custom2", "") or ""
                return ""

            for idx in range(total_slots):
                item = items[idx] if idx < len(items) else None
                # 시작 행 계산: 0번은 3행, 1번은 14행, 2번은 25행 ... (슬롯당 11행 불변)
                start_row = 3 + (idx * 11)
                
                # (1) 사진 영역 행 높이 설정
                # 전체 슬롯 11개 행의 합산 높이(347.25pt)는 절대 불변!
                if row_count == 1:
                    # 1줄 모드: 사진 영역이 9행 -> 10행으로 확장 (높이 324.75pt)
                    photo_end_row = start_row + 9
                    ws.row_dimensions[start_row].height = 17.25
                    for r in range(start_row + 1, start_row + 8):
                        ws.row_dimensions[r].height = 38.25
                    ws.row_dimensions[start_row + 8].height = 17.25
                    ws.row_dimensions[start_row + 9].height = 22.5  # 표에서 흡수한 높이로 사진 확장!
                    target_img_h = 404
                    box_h = 428
                else:
                    # 2줄 모드: 표준 9행 (높이 302.25pt)
                    photo_end_row = start_row + 8
                    ws.row_dimensions[start_row].height = 17.25
                    for r in range(start_row + 1, start_row + 8):
                        ws.row_dimensions[r].height = 38.25
                    ws.row_dimensions[start_row + 8].height = 17.25
                    target_img_h = 380
                    box_h = 403

                # 사진 테두리 적용 (A열 ~ I열 - 외곽 medium, 하단 구분선 thin)
                for r in range(start_row, photo_end_row + 1):
                    for c in range(1, 10):
                        cell = ws.cell(row=r, column=c)
                        t = medium_border_side if r == start_row else None
                        b = thin_border_side if r == photo_end_row else None
                        l = medium_border_side if c == 1 else None
                        rg = medium_border_side if c == 9 else None
                        cell.border = Border(top=t, bottom=b, left=l, right=rg)

                # (2) 이미지 처리 및 삽입
                if item is not None:
                    box_w = 655
                    target_img_w = 624
                    processed_img = process_picture(item, target_w=target_img_w, target_h=target_img_h)
                    
                    img_filename = f"img_{idx}_{item.item_id[:8]}.jpg"
                    temp_img_file = os.path.join(temp_dir, img_filename)
                    processed_img.save(temp_img_file, "JPEG", quality=92)
                    
                    xl_img = OpenpyxlImage(temp_img_file)
                    xl_img.width = target_img_w
                    xl_img.height = target_img_h

                    col_off_px = max(0, int((box_w - target_img_w) / 2))
                    row_off_px = max(0, int((box_h - target_img_h) / 2))
                    
                    marker = AnchorMarker(
                        col=0, 
                        colOff=pixels_to_EMU(col_off_px), 
                        row=start_row - 1, 
                        rowOff=pixels_to_EMU(row_off_px)
                    )
                    size = XDRPositiveSize2D(pixels_to_EMU(target_img_w), pixels_to_EMU(target_img_h))
                    xl_img.anchor = OneCellAnchor(_from=marker, ext=size)
                    ws.add_image(xl_img)

                # (3) 표 렌더링
                if row_count == 1:
                    # 단일행 표 (start_row + 10 행)
                    t_row = start_row + 10
                    ws.row_dimensions[t_row].height = 22.5
                    
                    if len(active_fields) == 1:
                        f = active_fields[0]
                        # A:B (라벨)
                        ws.merge_cells(start_row=t_row, start_column=1, end_row=t_row, end_column=2)
                        cell_lbl = ws.cell(row=t_row, column=1, value=f.get("label", ""))
                        cell_lbl.font = font_label
                        cell_lbl.alignment = align_center
                        # C:I (값)
                        ws.merge_cells(start_row=t_row, start_column=3, end_row=t_row, end_column=9)
                        cell_val = ws.cell(row=t_row, column=3, value=_get_val(item, f.get("key", "")))
                        cell_val.font = font_data
                        cell_val.alignment = align_center
                    else:
                        f1, f2 = active_fields[0], active_fields[1]
                        # A:B (라벨 1)
                        ws.merge_cells(start_row=t_row, start_column=1, end_row=t_row, end_column=2)
                        cell_lbl1 = ws.cell(row=t_row, column=1, value=f1.get("label", ""))
                        cell_lbl1.font = font_label
                        cell_lbl1.alignment = align_center
                        # C:D (값 1)
                        ws.merge_cells(start_row=t_row, start_column=3, end_row=t_row, end_column=4)
                        cell_val1 = ws.cell(row=t_row, column=3, value=_get_val(item, f1.get("key", "")))
                        cell_val1.font = font_data
                        cell_val1.alignment = align_center
                        # E:F (라벨 2)
                        ws.merge_cells(start_row=t_row, start_column=5, end_row=t_row, end_column=6)
                        cell_lbl2 = ws.cell(row=t_row, column=5, value=f2.get("label", ""))
                        cell_lbl2.font = font_label
                        cell_lbl2.alignment = align_center
                        # G:I (값 2)
                        ws.merge_cells(start_row=t_row, start_column=7, end_row=t_row, end_column=9)
                        cell_val2 = ws.cell(row=t_row, column=7, value=_get_val(item, f2.get("key", "")))
                        cell_val2.font = font_data
                        cell_val2.alignment = align_center

                    # 표 테두리 적용
                    for tc in range(1, 10):
                        l = medium_border_side if tc == 1 else thin_border_side
                        rg = medium_border_side if tc == 9 else thin_border_side
                        ws.cell(row=t_row, column=tc).border = Border(top=thin_border_side, bottom=medium_border_side, left=l, right=rg)

                else:
                    # 2줄 모드 (start_row + 9, start_row + 10)
                    t_row1 = start_row + 9
                    t_row2 = start_row + 10
                    ws.row_dimensions[t_row1].height = 22.5
                    ws.row_dimensions[t_row2].height = 22.5

                    if len(active_fields) == 3:
                        f1, f2, f3 = active_fields[0], active_fields[1], active_fields[2]
                        # 1행
                        ws.merge_cells(start_row=t_row1, start_column=1, end_row=t_row1, end_column=2)
                        c1 = ws.cell(row=t_row1, column=1, value=f1.get("label", ""))
                        c1.font = font_label
                        c1.alignment = align_center

                        ws.merge_cells(start_row=t_row1, start_column=3, end_row=t_row1, end_column=4)
                        c2 = ws.cell(row=t_row1, column=3, value=_get_val(item, f1.get("key", "")))
                        c2.font = font_data
                        c2.alignment = align_center

                        ws.merge_cells(start_row=t_row1, start_column=5, end_row=t_row1, end_column=6)
                        c3 = ws.cell(row=t_row1, column=5, value=f2.get("label", ""))
                        c3.font = font_label
                        c3.alignment = align_center

                        ws.merge_cells(start_row=t_row1, start_column=7, end_row=t_row1, end_column=9)
                        c4 = ws.cell(row=t_row1, column=7, value=_get_val(item, f2.get("key", "")))
                        c4.font = font_data
                        c4.alignment = align_center

                        # 2행 (f3 전폭 병합)
                        ws.merge_cells(start_row=t_row2, start_column=1, end_row=t_row2, end_column=2)
                        c5 = ws.cell(row=t_row2, column=1, value=f3.get("label", ""))
                        c5.font = font_label
                        c5.alignment = align_center

                        ws.merge_cells(start_row=t_row2, start_column=3, end_row=t_row2, end_column=9)
                        c6 = ws.cell(row=t_row2, column=3, value=_get_val(item, f3.get("key", "")))
                        c6.font = font_data
                        c6.alignment = align_center

                    else:
                        f1, f2, f3, f4 = active_fields[0], active_fields[1], active_fields[2], active_fields[3]
                        # 1행 (f1, f2)
                        ws.merge_cells(start_row=t_row1, start_column=1, end_row=t_row1, end_column=2)
                        c1 = ws.cell(row=t_row1, column=1, value=f1.get("label", "공    종"))
                        c1.font = font_label
                        c1.alignment = align_center

                        ws.merge_cells(start_row=t_row1, start_column=3, end_row=t_row1, end_column=4)
                        c2 = ws.cell(row=t_row1, column=3, value=_get_val(item, f1.get("key", "work_type")))
                        c2.font = font_data
                        c2.alignment = align_center

                        ws.merge_cells(start_row=t_row1, start_column=5, end_row=t_row1, end_column=6)
                        c3 = ws.cell(row=t_row1, column=5, value=f2.get("label", "내       용"))
                        c3.font = font_label
                        c3.alignment = align_center

                        ws.merge_cells(start_row=t_row1, start_column=7, end_row=t_row1, end_column=9)
                        c4 = ws.cell(row=t_row1, column=7, value=_get_val(item, f2.get("key", "content")))
                        c4.font = font_data
                        c4.alignment = align_center

                        # 2행 (f3, f4)
                        ws.merge_cells(start_row=t_row2, start_column=1, end_row=t_row2, end_column=2)
                        c5 = ws.cell(row=t_row2, column=1, value=f3.get("label", "위    치"))
                        c5.font = font_label
                        c5.alignment = align_center

                        ws.merge_cells(start_row=t_row2, start_column=3, end_row=t_row2, end_column=4)
                        c6 = ws.cell(row=t_row2, column=3, value=_get_val(item, f3.get("key", "location")))
                        c6.font = font_data
                        c6.alignment = align_center

                        ws.merge_cells(start_row=t_row2, start_column=5, end_row=t_row2, end_column=6)
                        c7 = ws.cell(row=t_row2, column=5, value=f4.get("label", "촬 영 일 자"))
                        c7.font = font_label
                        c7.alignment = align_center

                        ws.merge_cells(start_row=t_row2, start_column=7, end_row=t_row2, end_column=9)
                        c8 = ws.cell(row=t_row2, column=7, value=_get_val(item, f4.get("key", "date_str")))
                        c8.font = font_data
                        c8.alignment = align_center

                    # 2줄 테두리 적용
                    for tr in (t_row1, t_row2):
                        for tc in range(1, 10):
                            l = medium_border_side if tc == 1 else thin_border_side
                            rg = medium_border_side if tc == 9 else thin_border_side
                            t = thin_border_side
                            b = medium_border_side if tr == t_row2 else thin_border_side
                            ws.cell(row=tr, column=tc).border = Border(top=t, bottom=b, left=l, right=rg)

                # 2장 단위 자동 페이지 구분선 (Page Break) 추가
                if (idx + 1) % 2 == 0 and (idx + 1) < total_slots:
                    ws.row_breaks.append(Break(id=start_row + 10))

            # 저장
            wb.save(output_path)
            return True
        finally:
            if os.path.exists(temp_dir):
                shutil.rmtree(temp_dir, ignore_errors=True)
